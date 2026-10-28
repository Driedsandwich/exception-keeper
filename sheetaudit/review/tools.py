"""Bounded, read-only workbook lookups. Model-supplied arguments are untrusted."""
import builtins
from collections import Counter
import json

from ..context import cell_context
from openpyxl.utils.cell import coordinate_to_tuple
from ..llm.jsonout import validate
from ..r1c1 import normalize
from .schemas import obj, text
from .structure import range_bounds

TOOL_ARGS = {
    "read_range": obj({"sheet": text(), "range": text()}),
    "find_notes": obj({"sheet": text(), "row": {"type": "integer"}}),
    "find_rule": obj({"keyword": text(120)}),
    "same_pattern_cells": obj({"sheet": text(), "cell": text()}),
}
TOOLS = [{"type": "function", "function": {"name": name, "parameters": schema,
          "description": {"read_range": "Read at most 60 cells with formulas and saved values.",
                          "find_notes": "Read the notes column and its header for a row.",
                          "find_rule": "Find up to 10 text hits in sheets classified as rules.",
                          "same_pattern_cells": "Read up to 20 majority-pattern cells in the same formula run."}[name]}}
         for name, schema in TOOL_ARGS.items()]


def cell_record(cell):
    return dict(cell=f"{cell.sheet}!{cell.coordinate}", formula=cell.formula, value=cell.cached_value)


class WorkbookTools:
    def __init__(self, book, outline, candidate):
        self.book, self.outline, self.candidate = book, outline, candidate

    def execute(self, name, arguments):
        try:
            if name not in TOOL_ARGS:
                raise ValueError("Unknown tool")
            args = json.loads(arguments) if isinstance(arguments, str) else arguments
            if validate(args, TOOL_ARGS[name]):
                raise ValueError("Invalid tool arguments")
            return getattr(self, name)(**args)
        except (ValueError, TypeError, KeyError, OverflowError):
            return {"error": "Invalid tool arguments, unknown sheet, or range exceeds limits"}

    def read_range(self, sheet, range):
        tab, (r1, c1, r2, c2) = range_bounds(self.book, sheet, range, 60)
        return [cell_record(tab.at(r, c)) for r in builtins.range(r1, r2+1)
                for c in builtins.range(c1, c2+1)]

    def find_notes(self, sheet, row):
        tab = self.book.sheet(sheet)
        if not tab or not 1 <= row <= tab.max_row:
            raise ValueError("Unknown sheet or row")
        context = cell_context(self.book, tab.at(row, 1))
        address = context["notes_cell"]
        header = None
        if address:
            _, column = coordinate_to_tuple(address)
            from ..context import NOTES_HEADERS
            headers = [c for c in tab.cells.values() if c.column == column and c.row < row and not c.formula
                       and isinstance(c.cached_value, str) and c.cached_value.strip().casefold() in NOTES_HEADERS]
            if headers:
                header = cell_record(max(headers, key=lambda c: c.row))
        return dict(cell=f"{tab.name}!{address}" if address else None,
                    text=context["notes"], header=header)

    def find_rule(self, keyword):
        if not keyword.strip():
            raise ValueError("Empty keyword")
        names = [s["name"] for s in self.outline["sheets"] if s["role"] == "rules"]
        hits = []
        for name in names:
            for cell in self.book.sheets[name].cells.values():
                if (not cell.formula and isinstance(cell.cached_value, str)
                        and keyword.casefold() in cell.cached_value.casefold()):
                    hits.append(cell_record(cell))
                    if len(hits) == 10:
                        return hits
        return hits

    def same_pattern_cells(self, sheet, cell):
        tab, (r1, c1, r2, c2) = range_bounds(self.book, sheet, cell, 1)
        center = tab.at(r1, c1)
        result = {}
        for dr, dc in ((1, 0), (0, 1)):
            run = [center] if center.formula else []
            for sign in (-1, 1):
                row, col = r1 + sign*dr, c1 + sign*dc
                while 1 <= row <= tab.max_row and 1 <= col <= tab.max_column:
                    other = tab.at(row, col)
                    if not other.formula:
                        break
                    run.append(other)
                    row, col = row + sign*dr, col + sign*dc
            patterns = Counter(normalize(c.formula, c.coordinate) for c in run)
            if not patterns:
                continue
            majority, count = patterns.most_common(1)[0]
            if count * 2 <= len(run):
                continue
            for other in run:
                if other.coordinate != cell and normalize(other.formula, other.coordinate) == majority:
                    result[other.coordinate] = cell_record(other)
                    if len(result) == 20:
                        return list(result.values())
        return list(result.values())

    def upfront(self):
        c = self.candidate
        return {"read_range": self.execute("read_range", {"sheet": c.sheet, "range": c.cell}),
                "find_notes": self.execute("find_notes", {"sheet": c.sheet, "row": self.book.sheets[c.sheet].cells[c.cell].row}),
                "find_rule": self.execute("find_rule", {"keyword": c.context.get("row_label") or c.kind[0]}),
                "same_pattern_cells": self.execute("same_pattern_cells", {"sheet": c.sheet, "cell": c.cell})}
