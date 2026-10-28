"""Compact workbook discovery, followed by Nano's validated structural suggestions."""
import json
import re

from openpyxl.utils.cell import column_index_from_string

from ..context import NOTES_HEADERS
from ..llm.jsonout import request_json
from ..llm.models import NANO
from ..r1c1 import parse_reference
from .schemas import OUTLINE, SYSTEM


def compact_outline(book):
    rows = []
    for sheet in book.sheets.values():
        first_data = min((c.row for c in sheet.cells.values()
                          if c.formula or type(c.cached_value) in (int, float)), default=6)
        headers = [{"cell": c.coordinate, "text": c.cached_value[:100]}
                   for c in sheet.cells.values() if c.row < min(first_data, 6) and c.column <= 12
                   and isinstance(c.cached_value, str) and not c.formula][:30]
        labels = [{"cell": f"A{r}", "text": sheet.at(r, 1).cached_value[:80]}
                  for r in range(1, min(sheet.max_row, 10) + 1)
                  if isinstance(sheet.at(r, 1).cached_value, str)][:8]
        notes = [c.coordinate for c in sheet.cells.values() if c.row <= 5
                 and isinstance(c.cached_value, str) and not c.formula
                 and c.cached_value.strip().casefold() in NOTES_HEADERS][:8]
        totals = sorted({c.row for c in sheet.cells.values()
                         if c.formula and re.search(r"\b(?:SUM|AVERAGE|COUNT|MIN|MAX)\s*\(", c.formula, re.I)})[:40]
        rows.append(dict(name=sheet.name, dimensions=f"{sheet.max_row}x{sheet.max_column}",
                         headers=headers, labels=labels, notes_headers=notes, aggregate_rows=totals))
    return json.dumps({"sheets": rows}, ensure_ascii=False)[:6000]


def range_bounds(book, sheet_name, address, max_cells=None):
    sheet = book.sheet(sheet_name) if isinstance(sheet_name, str) else None
    ref = parse_reference(address) if isinstance(address, str) else None
    if not sheet or not ref or ref.sheet is not None:
        raise ValueError("Unknown sheet or invalid local A1 range")
    r1, c1, r2, c2 = ref.bounds
    if r2 > sheet.max_row or c2 > sheet.max_column:
        raise ValueError("Range outside workbook dimensions")
    if max_cells is not None and (r2-r1+1)*(c2-c1+1) > max_cells:
        raise ValueError(f"Range exceeds {max_cells} cells")
    return sheet, (r1, c1, r2, c2)


def validate_outline(value, book, warnings):
    result, seen = [], set()
    for item in value["sheets"]:
        sheet = book.sheet(item["name"])
        if not sheet or sheet.name in seen:
            warnings.append("Outline: unknown or duplicate sheet dropped")
            continue
        seen.add(sheet.name)
        entry = dict(item, name=sheet.name)
        row = entry.get("header_row")
        if row is not None and not 1 <= row <= sheet.max_row:
            entry.pop("header_row")
            warnings.append(f"Outline {sheet.name}: invalid header_row dropped")
        for key in ("label_column", "notes_column"):
            col = entry.get(key)
            if col is not None:
                try:
                    if not 1 <= column_index_from_string(col) <= sheet.max_column:
                        raise ValueError()
                    entry[key] = col.upper()
                except ValueError:
                    entry.pop(key)
                    warnings.append(f"Outline {sheet.name}: invalid {key} dropped")
        entry["total_rows"] = []
        for row in item["total_rows"]:
            if 1 <= row <= sheet.max_row:
                entry["total_rows"].append(row)
            else:
                warnings.append(f"Outline {sheet.name}: invalid total row dropped")
        entry["rules_ranges"] = []
        for address in item["rules_ranges"]:
            try:
                range_bounds(book, sheet.name, address)
                entry["rules_ranges"].append(address)
            except ValueError:
                warnings.append(f"Outline {sheet.name}: invalid rules range dropped")
        result.append(entry)
    return {"sheets": result}


def structure(book, call, warnings, schema_style=None):
    compact = compact_outline(book)
    messages = [{"role": "system", "content": SYSTEM},
                {"role": "user", "content": "Classify sheet roles and locate headers, notes, totals and rules. "
                 + "Columns must be A1 column letters. JSON schema: " + json.dumps(OUTLINE) + "\n" + compact}]
    value, errors = request_json(call, NANO, messages, OUTLINE, schema_style=schema_style, max_tokens=900)
    if value is None:
        warnings.append("Structure failed: " + "; ".join(errors))
        return {"sheets": []}
    return validate_outline(value, book, warnings)
