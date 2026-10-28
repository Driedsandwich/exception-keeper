"""Read formula text and saved values without modifying the workbook."""

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter


def is_number(value: object) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


@dataclass(frozen=True)
class Cell:
    sheet: str
    coordinate: str
    formula: str | None
    cached_value: object
    number_format: str = "General"
    comment: str | None = None

    @property
    def row(self) -> int:
        return coordinate_to_tuple(self.coordinate)[0]

    @property
    def column(self) -> int:
        return coordinate_to_tuple(self.coordinate)[1]

    @property
    def occupied(self) -> bool:
        return self.formula is not None or self.cached_value is not None


@dataclass
class Sheet:
    name: str
    cells: dict[str, Cell]
    max_row: int
    max_column: int

    def at(self, row: int, column: int) -> Cell:
        coordinate = f"{get_column_letter(column)}{row}"
        return self.cells.get(coordinate, Cell(self.name, coordinate, None, None))


@dataclass
class Workbook:
    sheets: dict[str, Sheet]

    def sheet(self, name: str) -> Sheet | None:
        # Excel sheet references are case-insensitive.
        return next((s for n, s in self.sheets.items() if n.casefold() == name.casefold()), None)


def load(path: str | Path) -> Workbook:
    formulas = load_workbook(path, data_only=False, keep_links=False)
    try:
        values = load_workbook(path, data_only=True, keep_links=False)
        try:
            sheets = {}
            for ws in formulas:
                if ws.max_row * ws.max_column > 200_000:
                    raise ValueError(f"Sheet {ws.title!r} exceeds the 200,000-cell scan limit")
                cells = {}
                for row in ws.iter_rows():
                    for cell in row:
                        if cell.value is None and cell.comment is None:
                            continue
                        formula = cell.value if cell.data_type == "f" else None
                        if formula is not None and not isinstance(formula, str):
                            raise ValueError(f"Unsupported array/data-table formula at {ws.title}!{cell.coordinate}")
                        cells[cell.coordinate] = Cell(
                            ws.title, cell.coordinate, formula,
                            values[ws.title][cell.coordinate].value,
                            cell.number_format, cell.comment.text if cell.comment else None,
                        )
                sheets[ws.title] = Sheet(ws.title, cells, ws.max_row, ws.max_column)
            return Workbook(sheets)
        finally:
            values.close()
    finally:
        formulas.close()
