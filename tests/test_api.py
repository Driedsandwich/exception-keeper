"""Fast, offline API, staged regression and preview transport checks."""
import base64
from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from datetime import datetime
from decimal import Decimal
from http.client import HTTPConnection
import io
import json
from pathlib import Path
import struct
import threading
from uuid import UUID, uuid4
import zipfile

from openpyxl import Workbook
import pytest

from sheetaudit.api import Dependencies, DISABLED, MAX_CANDIDATES, choose_lane, handle
from sheetaudit.budget import Budget, BudgetStopped, FileLedger, JST, MemoryLedger
from sheetaudit.llm.client import ChatResult, LLMError, MockChatClient
from sheetaudit.llm.models import NANO

ROOT = Path(__file__).resolve().parents[1]
NOW = lambda: datetime(2026, 10, 24, 12, tzinfo=JST)
ENV = {"EK_LIVE": "1", "EK_GRANTED_CREDITS_USD": "10", "NEBIUS_API_KEY": "fake-api-key",
       "EK_JUDGE_PASSPHRASE": "fake-judge-passphrase"}
NAMES = ["club_membership_fees", "small_business_budget", "event_budget"]


def script(name="event_budget"):
    return json.loads((ROOT / "tests/fixtures/mock_scripts" / (name + ".json")).read_text())


def setup(name="event_budget", ledger=None):
    return Dependencies(ledger if ledger is not None else MemoryLedger(), MockChatClient(script(name)), NOW)


def request(stage="start", name="event_budget", session=None, payload=None, **extra):
    result = {"stage": stage, "workbook_b64": base64.b64encode((ROOT / "samples" / (name + ".xlsx")).read_bytes()).decode()}
    if session is not None:
        result["session"] = session
    if payload is not None:
        result["payload"] = payload
    return dict(result, **extra)


def start(deps, name="event_budget", env=None, **extra):
    status, data = handle(request(name=name, **extra), ENV if env is None else env, deps)
    assert status == 200, data
    assert UUID(data["session"]).version == 4
    return data


def judge_request(data, index=0, name="event_budget", **extra):
    return request("judge", name, data["session"],
                   {"candidate_id": data["candidates"][index]["id"], "outline": data["outline"]}, **extra)


def run_flow(send, name):
    status, initial = send(request(name=name))
    assert status == 200, initial
    decisions, costs = [], [initial["session_cost_usd"]]
    for index, _ in enumerate(initial["candidates"][:initial["max_candidates"]]):
        status, data = send(judge_request(initial, index, name))
        assert status == 200, data
        decisions.append(data["decision"])
        costs.append(data["session_cost_usd"])
    status, final = send(request("summary", name, initial["session"], {"decisions": decisions, "outline": initial["outline"]}))
    assert status == 200, final
    assert costs == sorted(costs) and final["session_cost_usd"] > costs[-1]
    assert final["source"] == "mock"
    return final


@pytest.mark.parametrize("name", NAMES)
def test_split_flow_matches_precomputed(name):
    deps = setup(name)
    final = run_flow(lambda req: handle(req, ENV, deps), name)
    expected = json.loads((ROOT / "tests/fixtures/precomputed_mock" / (name + ".json")).read_text())
    # Exact finding equality includes verdicts, quotes, grounding, engine impact and fix.
    assert final["findings"] == expected["findings"]
    assert final["summary"] == expected["summary"]
    assert final["session_cost_usd"] == pytest.approx(expected["cost"]["total_usd"])
    for record in deps.ledger.records():
        assert {"app", "lane", "day", "amount", "session", "stage", "event"} <= record.keys()
        assert "workbook_b64" not in record and "outline" not in record and "decision" not in record


def test_mock_stages_are_independent_of_candidate_order_and_keep_tool_rounds():
    deps = setup()
    initial = start(deps)
    decisions = []
    for index in (1, 0):
        status, body = handle(judge_request(initial, index), ENV, deps)
        assert status == 200
        decisions.append(body["decision"])
    assert decisions[0]["verdict"] == "intended_exception"
    assert decisions[0]["grounding"] == "passed"
    assert len(deps.client.requests) == 4  # Nano, two Super tool rounds, other Super.


def test_keyed_script_repair_stays_with_candidate():
    fixture = script()
    answer = fixture["responses"][1]
    broken = dict(answer, content="not JSON")
    fixture["responses"].insert(1, broken)
    deps = Dependencies(MemoryLedger(), MockChatClient(fixture), NOW)
    initial = start(deps)
    status, body = handle(judge_request(initial), ENV, deps)
    assert status == 200 and body["decision"]["grounding"] == "passed"
    assert len(deps.client.requests) == 3  # Nano, malformed Super, repair.


@pytest.mark.parametrize("credits", [None, "", "0", "-1", "NaN", "Infinity", "invalid"])
def test_credits_not_configured(credits):
    env = {"EK_LIVE": "1"}
    if credits is not None:
        env["EK_GRANTED_CREDITS_USD"] = credits
    deps = setup()
    status, body = handle(request(), env, deps)
    assert status == 429 and body == {"error": "budget", "message": DISABLED}
    assert not deps.client.requests and not deps.ledger.records()


@pytest.mark.parametrize("stage", ["start", "judge", "summary"])
def test_live_disabled_at_every_stage(stage):
    deps = setup()
    initial = start(deps)
    status, judged = handle(judge_request(initial), ENV, deps)
    assert status == 200
    req = request() if stage == "start" else judge_request(initial, 1) if stage == "judge" else request(
        "summary", session=initial["session"], payload={"decisions": [judged["decision"]]})
    before = len(deps.client.requests)
    env = dict(ENV, EK_LIVE="0")
    status, body = handle(req, env, deps)
    assert status == 429 and body["message"] == DISABLED
    assert len(deps.client.requests) == before


@pytest.mark.parametrize("passphrase,lane", [(None, "public"), ("wrong", "public"), (ENV["EK_JUDGE_PASSPHRASE"], "judge")])
def test_passphrase_lane(passphrase, lane):
    deps = setup()
    initial = start(deps, **({"passphrase": passphrase} if passphrase is not None else {}))
    assert initial["lane"] == lane
    assert {r["lane"] for r in deps.ledger.records()} == {lane}
    assert choose_lane(passphrase, ENV) == lane
    status, _ = handle(judge_request(initial, **({"passphrase": passphrase} if passphrase is not None else {})), ENV, deps)
    assert status == 200


def test_lane_is_bound_to_session():
    deps = setup()
    initial = start(deps)
    status, body = handle(judge_request(initial, passphrase=ENV["EK_JUDGE_PASSPHRASE"]), ENV, deps)
    assert status == 400 and body["error"] == "bad_request"
    assert len(deps.client.requests) == 1


def test_judge_before_start_and_failed_start_are_refused():
    deps = setup()
    req = request("judge", session=str(uuid4()), payload={"candidate_id": "e9c3faca5778fcac"})
    assert handle(req, ENV, deps)[0] == 400
    session = req["session"]
    deps.ledger.add("ek", "public", "2026-10-24", 0, session=session, stage="start", event="started")
    deps.ledger.add("ek", "public", "2026-10-24", 0, session=session, stage="start", event="finished", success=False)
    assert handle(req, ENV, deps)[0] == 400


def test_summary_before_judge_duplicate_judge_and_post_summary_judge():
    deps = setup()
    initial = start(deps)
    req = request("summary", session=initial["session"], payload={"decisions": []})
    assert handle(req, ENV, deps)[0] == 400
    status, decision = handle(judge_request(initial), ENV, deps)
    assert status == 200
    assert handle(judge_request(initial), ENV, deps)[0] == 400
    req["payload"]["decisions"] = [decision["decision"]]
    assert handle(req, ENV, deps)[0] == 200
    assert handle(req, ENV, deps)[0] == 400  # One summary, including attempts.
    assert handle(judge_request(initial, 1), ENV, deps)[0] == 400


def test_inflight_stage_and_twelve_judge_attempts_fail_closed():
    deps = setup()
    initial = start(deps)
    for i in range(MAX_CANDIDATES):
        deps.ledger.add("ek", "public", "2026-10-24", 0, session=initial["session"], stage="judge", event="started", candidate_id=str(i))
        if i == 0:
            assert handle(judge_request(initial), ENV, deps)[0] == 400
        deps.ledger.add("ek", "public", "2026-10-24", 0, session=initial["session"], stage="judge", event="finished", success=False)
    assert handle(judge_request(initial), ENV, deps)[0] == 400
    assert len(deps.client.requests) == 1


@pytest.mark.parametrize("mutation", ["unknown_id", "duplicate", "impact", "verdict", "shape", "missing"])
def test_summary_rejects_untrusted_decisions(mutation):
    deps = setup()
    initial = start(deps)
    status, body = handle(judge_request(initial), ENV, deps)
    assert status == 200
    decisions = [deepcopy(body["decision"])]
    if mutation == "unknown_id": decisions[0]["candidate_id"] = "unknown"
    if mutation == "duplicate": decisions *= 2
    if mutation == "impact": decisions[0]["impact"] = 999999
    if mutation == "verdict": decisions[0]["verdict"] = "intended_exception"
    if mutation == "shape": decisions = ["not a decision"]
    if mutation == "missing": decisions = []
    status, result = handle(request("summary", session=initial["session"], payload={"decisions": decisions}), ENV, deps)
    assert status == 400 and result["error"] == "bad_request"


def test_upload_and_outline_are_bound_to_start():
    deps = setup()
    initial = start(deps)
    changed = judge_request(initial)
    changed["workbook_b64"] = request(name="club_membership_fees")["workbook_b64"]
    assert handle(changed, ENV, deps)[0] == 400
    changed = judge_request(initial)
    changed["payload"]["outline"] = {"sheets": []}
    assert handle(changed, ENV, deps)[0] == 400


@pytest.mark.parametrize("req", [None, [], {}, {"stage": []}, {"stage": "unknown"},
    {"stage": "start", "extra": 1}, {"stage": "start", "passphrase": 1},
    {"stage": "judge", "session": "invalid"}, {"stage": "start", "workbook_b64": "!!!"},
    {"stage": "start", "payload": []}, {"stage": "start", "session": str(uuid4())}])
def test_bad_requests(req):
    assert handle(req, ENV, setup())[0] == 400


def archive(members):
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in members:
            z.writestr(name, data)
    return out.getvalue()


@pytest.mark.parametrize("raw,code", [
    (b"not a zip", "unsupported_file"), (b"x" * 1_000_001, "too_large"),
    (archive([(str(i), "x") for i in range(201)]), "too_large"),
    (archive([("hello.txt", "x")]), "unsupported_file"),
])
def test_upload_refusals(raw, code):
    req = request(workbook_b64=base64.b64encode(raw).decode())
    status, body = handle(req, ENV, setup())
    assert body["error"] == code
    assert status == (413 if code == "too_large" else 415)


def test_zip_bomb_rejected_before_workbook_load(monkeypatch):
    # Inflate the directory's declared uncompressed size; allocate no large buffer.
    raw = bytearray(archive([("big.xml", "tiny")]))
    central = raw.index(b"PK\x01\x02")
    struct.pack_into("<I", raw, central + 24, 20_000_001)
    def forbidden(*args):
        pytest.fail("openpyxl must not see a rejected archive")
    monkeypatch.setattr("sheetaudit.api.load", forbidden)
    status, body = handle(request(workbook_b64=base64.b64encode(raw).decode()), ENV, setup())
    assert status == 413 and body["error"] == "too_large"


def test_cell_scan_limit_and_macro_package():
    book = Workbook()
    book.active["A200001"] = 1
    raw = io.BytesIO(); book.save(raw); book.close()
    assert handle(request(workbook_b64=base64.b64encode(raw.getvalue()).decode()), ENV, setup())[1]["error"] == "too_large"
    normal = base64.b64decode(request()["workbook_b64"])
    with zipfile.ZipFile(io.BytesIO(normal)) as z:
        members = [(n, z.read(n)) for n in z.namelist()]
    members.append(("xl/vbaProject.bin", b"fake"))
    assert handle(request(workbook_b64=base64.b64encode(archive(members)).decode()), ENV, setup())[1]["error"] == "unsupported_file"


@pytest.mark.parametrize("durable", [False, True])
def test_session_cap_from_ledger_across_handler_instances(tmp_path, durable):
    path = tmp_path / "ledger.jsonl"
    ledger = FileLedger(path) if durable else MemoryLedger()
    deps = setup(ledger=ledger)
    initial = start(deps)
    ledger.add("ek", "public", "2026-10-24", .1499, session=initial["session"], stage="judge", event="usage")
    separate = setup(ledger=FileLedger(path) if durable else ledger)
    status, body = handle(judge_request(initial), ENV, separate)
    assert status == 429 and body["error"] == "budget"
    assert not separate.client.requests
    # New sessions do not inherit another session's $0.15 cap.
    assert handle(request(), ENV, separate)[0] == 200


def test_granted_credits_count_both_apps_and_started_markers(tmp_path):
    ledger = FileLedger(tmp_path / "ledger.jsonl")
    ledger.add("pw", "judge", "2026-10-24", 0, event="started", session="pw-session", stage="extract")
    ledger.add("pw", "judge", "2026-10-24", .99999, event="finished", session="pw-session", stage="extract", success=True)
    deps = setup(ledger=ledger)
    status, body = handle(request(), dict(ENV, EK_GRANTED_CREDITS_USD="1"), deps)
    assert status == 429 and body["error"] == "budget" and not deps.client.requests
    assert ledger.total() == Decimal(".99999")
    assert ledger.spent_session("pw", "pw-session") == Decimal(".99999")


def test_budget_session_reconstructed_and_credits_cap():
    ledger = MemoryLedger()
    chat = MockChatClient([dict(content="{}", usage={"prompt_tokens": 10, "completion_tokens": 10})])
    first = Budget(ledger, now=NOW, session="one", credits_usd=1)
    first.complete(chat, "structure", "ek", "public", NANO, [], max_tokens=10)
    ledger.add("ek", "public", "2026-10-24", .15, session="one", stage="judge", event="usage")
    with pytest.raises(BudgetStopped, match="Per-session"):
        Budget(ledger, now=NOW, session="one", credits_usd=1).complete(chat, "judge", "ek", "public", NANO, [])


def test_parallel_same_candidate_admitted_once():
    deps = setup()
    initial = start(deps)
    with ThreadPoolExecutor(max_workers=2) as pool:
        statuses = list(pool.map(lambda _: handle(judge_request(initial), ENV, deps)[0], range(2)))
    assert sorted(statuses) == [200, 400]


def test_no_uploaded_workbook_written(tmp_path, monkeypatch):
    def forbidden(*args, **kwargs):
        pytest.fail("Uploaded bytes must never be written to a path")
    monkeypatch.setattr(Path, "write_bytes", forbidden)
    setup_deps = setup()
    run_flow(lambda req: handle(req, ENV, setup_deps), "event_budget")


def test_api_json_only_switch():
    deps = setup()
    initial = start(deps)
    status, body = handle(judge_request(initial), dict(ENV, EK_JUDGE_MODE="json_only"), deps)
    assert status == 200 and body["decision"]["verdict"] == "error"
    assert deps.client.requests[-1]["tools"] is None


def test_real_twelve_candidate_limit_and_not_reviewed_tail():
    from sheetaudit.engine import audit_book
    from sheetaudit.loader import load
    book = Workbook()
    sheet = book.active
    sheet.title = "Rates"
    for row in range(1, 14):
        sheet.cell(row, 1, "Tax rate")
        sheet.cell(row, 2, .1)
    stream = io.BytesIO(); book.save(stream); book.close()
    raw = stream.getvalue()
    candidates = audit_book(load(io.BytesIO(raw)))
    assert len(candidates) == 13
    responses = [dict(stage="structure", content='{"sheets": []}')]
    for candidate in candidates[:12]:
        answer = dict(candidate_id=candidate.id, verdict="assumption_to_verify", confidence="low",
                      reason="Verify the rate.", evidence=[dict(cell="Rates!" + candidate.cell, quote="0.1")],
                      fix=None, question="Is the rate current?")
        responses.append(dict(stage="judge", candidate_id=candidate.id, content=json.dumps(answer)))
    responses.append(dict(stage="summary", content=json.dumps(dict(headline="Check rates.",
                     order=[c.id for c in candidates], conflicts=[], question_for_user="Are the rates current?"))))
    deps = Dependencies(MemoryLedger(), MockChatClient({"responses": responses}), NOW)
    encoded = base64.b64encode(raw).decode()
    def send(stage, session=None, payload=None):
        return handle(dict(stage=stage, workbook_b64=encoded, session=session, payload=payload or {}), ENV, deps)
    status, initial = send("start")
    assert status == 200 and initial["max_candidates"] == 12
    decisions = []
    for candidate in candidates[:12]:
        status, result = send("judge", initial["session"], {"candidate_id": candidate.id})
        assert status == 200
        decisions.append(result["decision"])
    assert send("judge", initial["session"], {"candidate_id": candidates[-1].id})[0] == 400
    status, final = send("summary", initial["session"], {"decisions": decisions})
    assert status == 200 and len(final["findings"]) == 13
    assert final["findings"][-1]["verdict"] == "not reviewed (limit)"
    assert candidates[-1].id not in final["summary"]["order"]


def test_http_transport_without_socket():
    from sheetaudit.http import post
    deps = setup()
    class Transport:
        path = "/api/ek"
        def __init__(self, req):
            raw = json.dumps(req).encode()
            self.headers = {"Content-Length": str(len(raw))}
            self.rfile, self.wfile = io.BytesIO(raw), io.BytesIO()
        def send_response(self, status): self.status = status
        def send_header(self, *args): pass
        def end_headers(self): pass
    def send(req):
        transport = Transport(req)
        post(transport, deps, ENV)
        return transport.status, json.loads(transport.wfile.getvalue())
    result = run_flow(send, "event_budget")
    assert len(result["findings"]) == 2


@pytest.mark.parametrize("failure", [False, True])
def test_secrets_never_in_responses_logs_errors_or_ledger(failure, monkeypatch, capsys, caplog):
    for key, value in ENV.items():
        monkeypatch.setenv(key, value)
    fixture = script()
    leaked = ENV["NEBIUS_API_KEY"] + " " + ENV["EK_JUDGE_PASSPHRASE"]
    if failure:
        fixture["responses"][0] = dict(fixture["responses"][0], error=leaked)
    else:
        item = fixture["responses"][1]
        answer = json.loads(item["content"]); answer["reason"] = leaked
        item["content"] = json.dumps(answer)
    deps = Dependencies(MemoryLedger(), MockChatClient(fixture), NOW)
    if failure:
        status, body = handle(request(passphrase=ENV["EK_JUDGE_PASSPHRASE"]), ENV, deps)
        assert status == 502
    else:
        initial = start(deps, passphrase=ENV["EK_JUDGE_PASSPHRASE"])
        status, body = handle(judge_request(initial, passphrase=ENV["EK_JUDGE_PASSPHRASE"]), ENV, deps)
        assert status == 200
    captured = capsys.readouterr()
    all_output = (json.dumps(body) + json.dumps(deps.ledger.records()) + captured.out + captured.err + caplog.text
                  + str(LLMError(500, leaked)))
    for key in ("NEBIUS_API_KEY", "EK_JUDGE_PASSPHRASE"):
        assert ENV[key] not in all_output
    assert leaked not in str(LLMError(500, leaked))


def test_local_server_full_sample_with_mock_client(capsys, caplog):
    from serve import make_server
    server = make_server(("127.0.0.1", 0), deps=setup(), env=ENV, bind_and_activate=False)
    try:
        server.server_bind()
    except OSError as exc:
        server.server_close()
        pytest.skip(f"Loopback bind failed: {exc}")
    server.server_activate()
    thread = threading.Thread(target=server.serve_forever, kwargs={"poll_interval": .01}, daemon=True)
    thread.start()
    try:
        connection = HTTPConnection("127.0.0.1", server.server_port, timeout=2)
        def send(req):
            connection.request("POST", "/api/ek", json.dumps(req), {"Content-Type": "application/json"})
            response = connection.getresponse()
            return response.status, json.loads(response.read())
        final = run_flow(send, "event_budget")
        assert len(final["findings"]) == 2
        for path in ("/", "/app.mjs", "/samples/precomputed/event_budget.json", "/samples/event_budget.xlsx"):
            connection.request("GET", path)
            response = connection.getresponse(); assert response.status == 200; response.read()
        connection.close()
    finally:
        server.shutdown(); thread.join(timeout=1); server.server_close()
    captured = capsys.readouterr()
    for key in ("NEBIUS_API_KEY", "EK_JUDGE_PASSPHRASE"):
        assert ENV[key] not in captured.out + captured.err + caplog.text


@pytest.mark.parametrize("env_style,expected", [(None, "json_schema"), ("json_object", "json_object")])
def test_live_default_schema_style_is_strict_schema(env_style, expected):
    deps = setup()
    env = dict(ENV, **({"EK_SCHEMA_STYLE": env_style} if env_style else {}))
    start(deps, env=env)
    assert deps.client.requests[0]["response_format"]["type"] == expected
